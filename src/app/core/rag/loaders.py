"""
Lectura y troceado de documentos.

Mismos parámetros de troceado que la ingesta original por línea de comandos
(1000 caracteres con 200 de traslape), para que un documento subido desde el
panel se comporte igual que los que ya están indexados.
"""
import logging
from pathlib import Path
from typing import List

logger = logging.getLogger(__name__)

# .xls (formato viejo de Excel) queda fuera: openpyxl no lo lee y aceptarlo
# solo llevaría a un error al indexar.
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".xlsx")


def _pdf_reader(file_handle):
    """
    Lector de PDF.

    El proyecto declara pypdf en requirements.txt; PyPDF2 es su antecesor y solo
    está presente en entornos viejos. Se intenta el primero y se cae al segundo.
    """
    try:
        from pypdf import PdfReader
    except ImportError:  # entornos anteriores a la migración
        from PyPDF2 import PdfReader
    return PdfReader(file_handle)


def read_pdf(file_path: Path) -> str:
    text = ""
    try:
        with open(file_path, "rb") as f:
            reader = _pdf_reader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception as e:
        logger.error(f"Error leyendo PDF {file_path.name}: {e}")
    return text


def read_docx(file_path: Path) -> str:
    from docx import Document

    try:
        doc = Document(str(file_path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        logger.error(f"Error leyendo DOCX {file_path.name}: {e}")
        return ""


def read_xlsx(file_path: Path) -> str:
    """Convierte cada fila en 'Columna: valor | Columna: valor' para búsqueda semántica."""
    import openpyxl

    rows_text: List[str] = []
    try:
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        for sheet in wb.worksheets:
            headers: List[str] = []
            first_row = True
            for row in sheet.iter_rows(values_only=True):
                if all(v is None for v in row):
                    continue
                if first_row:
                    headers = [
                        str(h).strip() if h is not None else f"Col{i}"
                        for i, h in enumerate(row)
                    ]
                    first_row = False
                    continue
                parts = [
                    f"{h}: {str(v).strip()}"
                    for h, v in zip(headers, row, strict=False)
                    if v is not None and str(v).strip()
                ]
                if parts:
                    rows_text.append(" | ".join(parts))
        wb.close()
    except Exception as e:
        logger.error(f"Error leyendo XLSX {file_path.name}: {e}")
    return "\n".join(rows_text)


def read_txt(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        logger.error(f"Error leyendo texto {file_path.name}: {e}")
        return ""


def read_file(file_path: Path) -> str:
    """Elige el lector según la extensión."""
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return read_pdf(file_path)
    if ext == ".docx":
        return read_docx(file_path)
    if ext == ".xlsx":
        return read_xlsx(file_path)
    if ext in (".txt", ".md"):
        return read_txt(file_path)
    logger.warning(f"Tipo de archivo no soportado: {ext} → {file_path.name}")
    return ""


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> List[str]:
    """Divide el texto en fragmentos con traslape, descartando los vacíos."""
    if not text:
        return []
    if overlap >= chunk_size:
        raise ValueError("El traslape debe ser menor que el tamaño del fragmento.")

    chunks: List[str] = []
    start = 0
    text_len = len(text)
    while start < text_len:
        end = start + chunk_size
        fragment = text[start:end].strip()
        if fragment:
            chunks.append(fragment)
        start = end - overlap
    return chunks
